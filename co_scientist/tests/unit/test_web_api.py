"""JSON API the React app calls, served from the database.

Upload and research tests run on both backends (`api_cfg`); the Atlas leg needs
CO_SCI_ATLAS_TESTS=1. The Supervisor run itself is replaced with a no-op.
"""

from __future__ import annotations

import io
import re
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

import httpx
import pytest
import pytest_asyncio
from structlog.testing import capture_logs

from co_scientist import ids
from co_scientist.agents.supervisor import Supervisor
from co_scientist.config import Config
from co_scientist.models import Hypothesis, ResearchPlan, Session
from co_scientist.models.hypothesis import CitedPaper
from co_scientist.storage import db as db_mod
from co_scientist.storage.repos import analysis as analysis_repo
from co_scientist.storage.repos import attachments as att_repo
from co_scientist.storage.repos import blobs
from co_scientist.storage.repos import hypotheses as hyp_repo
from co_scientist.storage.repos import sessions as sess_repo
from co_scientist.storage.repos import syntheses as syn_repo
from co_scientist.tests import atlas_support
from co_scientist.web.app import create_app

from .test_research_view import (
    DATA_ANALYSIS_KEYS,
    DATA_POINT_KEYS,
    DATA_SUMMARY_KEYS,
    HYPOTHESIS_KEYS,
    RESEARCH_SESSION_KEYS,
    SOURCE_KEYS,
)

KEY = "EGFR|mutation|ERLOTINIB"
PNG = b"\x89PNG\r\n\x1a\nfake"


# ----------------------------------------------------------------- fixtures --


@pytest_asyncio.fixture(params=["sqlite", "atlas"])
async def api_cfg(request: pytest.FixtureRequest, tmp_cfg: Config, tmp_path: Path):
    """App config on each backend; the atlas leg uses a throwaway collection prefix."""
    if request.param == "sqlite":
        yield tmp_cfg
        return
    if not atlas_support.ENABLED:
        pytest.skip("set CO_SCI_ATLAS_TESTS=1 to run live Atlas tests")
    cfg = atlas_support.make_atlas_cfg()
    cfg.storage.data_dir = str(tmp_path)
    from co_scientist.storage.atlas.indexes import ensure_indexes

    conn = await db_mod.connect(cfg)
    try:  # B-tree only: search index builds cost minutes and these tests don't search
        await ensure_indexes(conn, cfg, search=False)
    finally:
        await conn.close()
    try:
        yield cfg
    finally:
        conn = await db_mod.connect(cfg)
        try:
            await atlas_support.drop_prefixed(conn.db, conn.prefix)
        finally:
            await conn.close()


@pytest.fixture
def started(monkeypatch: pytest.MonkeyPatch) -> list[dict[str, Any]]:
    calls: list[dict[str, Any]] = []

    async def fake_run(self, goal, **kw):
        calls.append({"goal": goal, "budget_usd": self.cfg.run.budget_usd, **kw})
        return kw.get("session_id")

    monkeypatch.setattr(Supervisor, "run_session", fake_run)
    return calls


def _client(cfg: Config) -> httpx.AsyncClient:
    return httpx.AsyncClient(
        transport=httpx.ASGITransport(app=create_app(cfg)), base_url="http://test"
    )


async def _make_session(cfg: Config, *, final_overview: str | None = None) -> str:
    conn = await db_mod.connect(cfg)
    try:
        sid = ids.session_id()
        now = datetime.now(UTC)
        await sess_repo.insert(conn, Session(
            id=sid, created_at=now, updated_at=now, status="running",
            research_goal="g", research_plan=ResearchPlan(objective="g"),
            config_snapshot={}, budget_tokens=1_000_000, budget_usd=10.0,
        ))
        if final_overview:
            await sess_repo.set_final_overview(conn, sid, final_overview)
        return sid
    finally:
        await conn.close()


async def _add_hyp(conn, sid: str, title: str, *, state: str, elo: float | None,
                   key: str | None = None) -> None:
    await hyp_repo.insert(conn, Hypothesis(
        id=ids.hypothesis_id(sid, "generation", title), session_id=sid,
        created_at=datetime.now(UTC), created_by="generation", strategy="data_driven",
        title=title, summary=f"{title} summary.",
        full_text=f"# {title}\n\n**Hypothesis.** {title} holds.\n\n## Mechanism\nWhy.",
        citations=[
            CitedPaper(title="A review of EGFR", url="https://www.nature.com/r", year=2020,
                       excerpt="EGFR matters"),
            CitedPaper(title="notes.pdf", url="upload://att_X#p1"),
        ],
        data_evidence={"candidate_key": key, "biomarker": "EGFR",
                       "alteration_type": "mutation", "drug": "ERLOTINIB",
                       "p_value": 0.01, "sample_size": 40} if key else None,
        artifact_path="x.json", elo=elo, state=state,
    ))


def _docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Results", 1)
    d.add_paragraph("EGFR mutations predict response to erlotinib in this cohort.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------- POST /api/sessions --


async def test_post_sessions_stores_files_and_returns_id(api_cfg, started) -> None:
    files = [
        ("files", ("clinical.csv", b"a,b\n1,2\n", "text/csv")),
        ("files", ("paper.docx", _docx(), "application/octet-stream")),
    ]
    async with _client(api_cfg) as c:
        r = await c.post("/api/sessions", data={"question": "Does EGFR matter?",
                                                "budget_usd": "2.5"}, files=files)
        assert r.status_code == 201, r.text
        sid = r.json()["session_id"]
        assert re.fullmatch(r"ses_[0-9A-Z]{26}", sid)
        listed = (await c.get(f"/api/sessions/{sid}/attachments")).json()["attachments"]
    by_name = {a["filename"]: a for a in listed}
    assert set(by_name) == {"clinical.csv", "paper.docx"}
    assert all(set(a) == {"id", "filename", "kind", "size", "status"} for a in listed)
    assert by_name["clinical.csv"]["kind"] == "tabular"
    assert by_name["clinical.csv"]["status"] == "stored"
    # Documents are ingested into the library on Atlas only.
    want = "ingested" if api_cfg.storage.backend == "atlas" else "stored"
    assert by_name["paper.docx"]["status"] == want

    conn = await db_mod.connect(api_cfg)
    try:
        session = await sess_repo.fetch(conn, sid)
        assert session is not None and session.budget_usd == 2.5
        metas = await att_repo.list_for_session(conn, sid)
        csv_meta = next(m for m in metas if m.filename == "clinical.csv")
        assert await blobs.get(conn, "attachments", csv_meta.blob_id) == b"a,b\n1,2\n"
    finally:
        await conn.close()
    (call,) = started
    assert call["session_id"] == sid and call["goal"] == "Does EGFR matter?"
    assert call["budget_usd"] == 2.5
    assert not call.get("attachments")
    assert api_cfg.run.budget_usd != 2.5  # per-session copy, app config untouched


async def test_post_sessions_rejects_oversize_and_bad_ext(tmp_cfg, started) -> None:
    tmp_cfg.attachments.max_file_mb = 1
    async with _client(tmp_cfg) as c:
        bad = await c.post("/api/sessions", data={"question": "q"},
                           files=[("files", ("evil.exe", b"MZ", "application/octet-stream"))])
        assert bad.status_code == 400
        assert "unsupported file type" in bad.json()["detail"]
        big = await c.post("/api/sessions", data={"question": "q"},
                           files=[("files", ("big.csv", b"x" * (1024 * 1024 + 1), "text/csv"))])
        assert big.status_code == 400
        assert "larger than 1 MB" in big.json()["detail"]
        tmp_cfg.attachments.max_unzipped_mb = 0
        unzipped = await c.post("/api/sessions", data={"question": "q"},
                                files=[("files", ("a.csv", b"a,b\n", "text/csv"))])
        assert unzipped.status_code == 413
        empty = await c.post("/api/sessions", data={"question": "  "})
        assert empty.status_code == 400
    assert started == []
    conn = await db_mod.connect(tmp_cfg)
    try:
        assert await sess_repo.list_with_stats(conn) == []
    finally:
        await conn.close()


async def test_post_sessions_enforces_session_and_total_quota(tmp_cfg, started) -> None:
    async with _client(tmp_cfg) as c:
        tmp_cfg.attachments.max_session_mb = 0
        r = await c.post("/api/sessions", data={"question": "q"},
                         files=[("files", ("a.csv", b"a,b\n", "text/csv"))])
        assert r.status_code == 413 and "per-session" in r.json()["detail"]
        tmp_cfg.attachments.max_session_mb = 125
        tmp_cfg.attachments.max_total_mb = 0
        r = await c.post("/api/sessions", data={"question": "q"},
                         files=[("files", ("a.csv", b"a,b\n", "text/csv"))])
        assert r.status_code == 413 and "storage limit" in r.json()["detail"]
    assert started == []
    conn = await db_mod.connect(tmp_cfg)
    try:
        assert await sess_repo.list_with_stats(conn) == []  # no session row on a 413
    finally:
        await conn.close()


async def test_post_sessions_ingest_failure_does_not_fail_upload(
    tmp_cfg, started, monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Force the Atlas branch on a SQLite connection and make ingestion blow up.
    import co_scientist.agents.supervisor as sup_mod

    calls: list[str] = []

    async def boom(conn, cfg, meta):
        calls.append(meta.filename)
        raise RuntimeError("secret detail")

    monkeypatch.setattr(sup_mod, "_is_atlas", lambda conn: True)
    monkeypatch.setattr(sup_mod, "ingest_attachment", boom)
    with capture_logs() as logs:
        async with _client(tmp_cfg) as c:
            r = await c.post("/api/sessions", data={"question": "q"}, files=[
                ("files", ("paper.pdf", b"%PDF-1.4 x", "application/pdf")),
                ("files", ("a.csv", b"a,b\n", "text/csv")),
            ])
    assert r.status_code == 201
    assert calls == ["paper.pdf"]  # tabular files are never ingested
    assert "secret detail" not in r.text
    assert "secret detail" not in repr(logs)
    failed = [e for e in logs if e["event"] == "upload_ingest_failed"]
    assert failed and failed[0]["error"] == "RuntimeError"


async def test_background_run_db_error_logs_type_only(
    tmp_cfg, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import asyncio

    from pymongo.errors import ServerSelectionTimeoutError

    async def boom(self, goal, **kw):
        raise ServerSelectionTimeoutError("secret host db0.internal:27017")

    monkeypatch.setattr(Supervisor, "run_session", boom)
    app = create_app(tmp_cfg)
    transport = httpx.ASGITransport(app=app)
    with capture_logs() as logs:
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
            r = await c.post("/api/sessions", data={"question": "q"})
            assert r.status_code == 201
            task = app.state.background_runs.get(r.json()["session_id"])
            if task is not None:
                await task
            # The older HTML form route starts its run the same way.
            form = await c.post("/sessions/new", data={"goal": "g"})
            assert form.status_code == 303
            for _ in range(20):
                await asyncio.sleep(0)
    failures = [e for e in logs if e["event"] == "background_run_failed"]
    assert len(failures) == 2
    assert all(e["err_type"] == "ServerSelectionTimeoutError" for e in failures)
    assert all("exc_info" not in e for e in failures)  # no traceback rendered
    assert "secret host" not in repr(logs)


async def test_prepare_session_marks_half_created_session_failed(
    tmp_cfg, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import co_scientist.agents.supervisor as sup_mod
    from co_scientist.attachments.intake import IncomingFile

    async def boom(conn, cfg, session_id, files):
        raise RuntimeError("store failed")

    monkeypatch.setattr(sup_mod, "store_incoming", boom)
    conn = await db_mod.connect(tmp_cfg)
    try:
        with pytest.raises(RuntimeError, match="store failed"):
            await Supervisor(tmp_cfg).prepare_session(
                conn, "g", attachments=[IncomingFile("a.csv", b"a\n", "text/csv", "tabular")],
            )
        (row,) = await sess_repo.list_with_stats(conn)
        session = await sess_repo.fetch(conn, row["id"])
        assert session is not None and session.status == "failed"
    finally:
        await conn.close()


# ---------------------------------------------------------------- /research --


async def test_research_matches_frontend_shape(api_cfg) -> None:
    sid = await _make_session(api_cfg)
    conn = await db_mod.connect(api_cfg)
    try:
        await _add_hyp(conn, sid, "alpha", state="in_tournament", elo=1250.0, key=KEY)
        await _add_hyp(conn, sid, "beta", state="pinned", elo=1300.0)
        await analysis_repo.save_report(conn, sid, {"candidates": [{"candidate_key": KEY}]})
        await analysis_repo.save_points(conn, sid, KEY, metric="os_months", unit="months",
                                        points=[{"sample_id": "S1", "group": "altered",
                                                 "value": 4.0}])
    finally:
        await conn.close()
    async with _client(api_cfg) as c:
        r = await c.get(f"/api/sessions/{sid}/research")
    assert r.status_code == 200, r.text
    body = r.json()
    assert set(body) == RESEARCH_SESSION_KEYS
    assert body["id"] == sid and body["isDemo"] is False
    assert [h["title"] for h in body["hypotheses"]] == ["beta", "alpha"]
    assert [h["rank"] for h in body["hypotheses"]] == [1, 2]
    for h in body["hypotheses"]:
        assert set(h) == HYPOTHESIS_KEYS
        assert set(h["dataSummary"]) == DATA_SUMMARY_KEYS
        assert all(set(p) == DATA_POINT_KEYS for p in h["dataPoints"])
        assert all(set(s) == SOURCE_KEYS for s in h["sources"])
    alpha = body["hypotheses"][1]
    assert alpha["dataPoints"] == [
        {"id": "S1", "label": "S1", "group": "altered", "value": 4.0, "event": None},
    ]
    assert set(alpha["dataAnalysis"]) == DATA_ANALYSIS_KEYS
    assert alpha["dataSummary"]["metric"] == "os_months"
    assert alpha["abstract"] == "**Hypothesis.** alpha holds."
    assert alpha["sources"][0]["evidenceType"] == "Review"
    assert alpha["sources"][0]["publication"] == "www.nature.com"
    assert alpha["sources"][1]["url"] is None
    assert body["hypotheses"][0]["dataPoints"] == []


async def test_research_excludes_rejected_and_quarantined(api_cfg) -> None:
    sid = await _make_session(api_cfg)
    conn = await db_mod.connect(api_cfg)
    try:
        await _add_hyp(conn, sid, "keep", state="in_tournament", elo=1200.0)
        for state in ("rejected", "quarantined", "retired", "draft"):
            await _add_hyp(conn, sid, state, state=state, elo=1500.0)
        await _add_hyp(conn, sid, "reviewed", state="reviewed", elo=None)
    finally:
        await conn.close()
    async with _client(api_cfg) as c:
        body = (await c.get(f"/api/sessions/{sid}/research")).json()
    assert [h["title"] for h in body["hypotheses"]] == ["keep"]


async def test_research_includes_drafts_and_reviewed(api_cfg) -> None:
    sid = await _make_session(api_cfg)
    conn = await db_mod.connect(api_cfg)
    try:
        await _add_hyp(conn, sid, "keep", state="in_tournament", elo=1200.0)
        await _add_hyp(conn, sid, "draft_hyp", state="draft", elo=None)
        await _add_hyp(conn, sid, "reviewed_hyp", state="reviewed", elo=1100.0)
        await _add_hyp(conn, sid, "rejected_hyp", state="rejected", elo=1300.0)
    finally:
        await conn.close()
    async with _client(api_cfg) as c:
        body = (await c.get(f"/api/sessions/{sid}/research?include_drafts=true")).json()
    titles = [h["title"] for h in body["hypotheses"]]
    assert "keep" in titles
    assert "draft_hyp" in titles
    assert "reviewed_hyp" in titles
    assert "rejected_hyp" not in titles
    hyp_map = {h["title"]: h for h in body["hypotheses"]}
    assert hyp_map["draft_hyp"]["state"] == "draft"
    assert hyp_map["reviewed_hyp"]["state"] == "reviewed"


async def test_research_loads_only_linked_points(
    api_cfg, monkeypatch: pytest.MonkeyPatch,
) -> None:
    sid = await _make_session(api_cfg)
    conn = await db_mod.connect(api_cfg)
    try:
        await _add_hyp(conn, sid, "linked", state="in_tournament", elo=1200.0, key=KEY)
        await _add_hyp(conn, sid, "gone", state="rejected", elo=1200.0, key="KRAS|mutation|X")
        await _add_hyp(conn, sid, "plain", state="in_tournament", elo=1100.0)
    finally:
        await conn.close()
    seen: list[list[str]] = []
    real = analysis_repo.fetch_points

    async def spy(conn, session_id, candidate_keys):
        seen.append(list(candidate_keys))
        return await real(conn, session_id, candidate_keys)

    monkeypatch.setattr(analysis_repo, "fetch_points", spy)
    async with _client(api_cfg) as c:
        assert (await c.get(f"/api/sessions/{sid}/research")).status_code == 200
    assert seen == [[KEY]]


async def test_research_404_for_unknown_session(tmp_cfg) -> None:
    async with _client(tmp_cfg) as c:
        r = await c.get(f"/api/sessions/{ids.session_id()}/research")
    assert r.status_code == 404


# -------------------------------------------------- synthesis/overview/sources --


async def test_synthesis_routes_read_db(tmp_cfg) -> None:
    sid = await _make_session(tmp_cfg)
    doc = {
        "overview": {"markdown": "# Overview"},
        "sources": {"session_id": sid, "generated_at": "t", "total_sources": 1,
                    "sources": [{"id": "S1"}]},
        "sources_list": [{"id": "S1"}],
    }
    conn = await db_mod.connect(tmp_cfg)
    try:
        await syn_repo.save(conn, sid, doc)
    finally:
        await conn.close()
    other = await _make_session(tmp_cfg)
    async with _client(tmp_cfg) as c:
        assert (await c.get(f"/api/sessions/{sid}/synthesis")).json() == doc
        assert (await c.get(f"/api/sessions/{sid}/overview")).json() == {"markdown": "# Overview"}
        assert (await c.get(f"/api/sessions/{sid}/sources")).json() == doc["sources"]
        for route in ("synthesis", "overview", "sources"):
            assert (await c.get(f"/api/sessions/{other}/{route}")).status_code == 404


async def test_synthesis_routes_fall_back_to_legacy_file(tmp_cfg) -> None:
    import json

    sid = await _make_session(tmp_cfg)
    final = tmp_cfg.data_dir / "artifacts" / sid / "final"
    final.mkdir(parents=True)
    (final / "synthesis.json").write_text(json.dumps(
        {"overview": {"markdown": "old"}, "sources": {"sources": []}}
    ))
    async with _client(tmp_cfg) as c:
        assert (await c.get(f"/api/sessions/{sid}/overview")).json() == {"markdown": "old"}
        assert (await c.get(f"/api/sessions/{sid}/sources")).json() == {"sources": []}


# ------------------------------------------------- timeline/analysis/plots --


async def test_timeline_route_pages_after_seq(tmp_cfg) -> None:
    from co_scientist import timeline

    sid = await _make_session(tmp_cfg)
    conn = await db_mod.connect(tmp_cfg)
    try:
        for i in range(3):
            await timeline.post(conn, sid, kind="note", title=f"m{i}")
    finally:
        await conn.close()
    async with _client(tmp_cfg) as c:
        all_msgs = (await c.get(f"/api/sessions/{sid}/timeline")).json()["messages"]
        later = (await c.get(f"/api/sessions/{sid}/timeline?after_seq=2")).json()["messages"]
        assert (await c.get(f"/api/sessions/{sid}/timeline?after_seq=-1")).status_code == 422
    assert [m["title"] for m in all_msgs] == ["m0", "m1", "m2"]
    assert set(all_msgs[0]) == {"seq", "ts", "role", "agent", "kind", "title",
                                "body_markdown", "refs"}
    assert [m["title"] for m in later] == ["m2"]


async def _seed_report(cfg: Config) -> str:
    sid = await _make_session(cfg)
    conn = await db_mod.connect(cfg)
    try:
        blob_id = await blobs.put(conn, "plots", "kaplan_meier_1.png", PNG, {"session_id": sid})
        await analysis_repo.save_report(conn, sid, {
            "summary_markdown": "s", "tests": [], "methods": {}, "candidates": [],
            "selection": {}, "sample_count": 3, "drugs": ["D"],
            "plots": [{"name": "kaplan_meier_1", "blob_id": blob_id,
                       "content_type": "image/png"}],
        })
    finally:
        await conn.close()
    return sid


async def test_analysis_route_hides_blob_ids(tmp_cfg) -> None:
    sid = await _seed_report(tmp_cfg)
    async with _client(tmp_cfg) as c:
        r = await c.get(f"/api/sessions/{sid}/analysis")
        missing = await c.get(f"/api/sessions/{await _make_session(tmp_cfg)}/analysis")
    assert r.status_code == 200
    body = r.json()
    assert body["sample_count"] == 3
    assert body["plots"] == [{"name": "kaplan_meier_1", "content_type": "image/png",
                              "url": f"/api/sessions/{sid}/plots/kaplan_meier_1"}]
    assert "blob_id" not in r.text
    assert missing.status_code == 404


async def test_analysis_route_serves_points_and_survival_curves(tmp_cfg) -> None:
    sid = await _seed_report(tmp_cfg)
    conn = await db_mod.connect(tmp_cfg)
    try:
        await analysis_repo.save_points(
            conn, sid, "TP53|mutation|CISPLATIN",
            metric="os_months", unit="months",
            points=[{"sample_id": "S1", "group": "altered", "value": 4.5}],
        )
    finally:
        await conn.close()
    async with _client(tmp_cfg) as c:
        r = await c.get(f"/api/sessions/{sid}/analysis")
    assert r.status_code == 200
    body = r.json()
    assert "points" in body
    assert "TP53|mutation|CISPLATIN" in body["points"]
    assert body["points"]["TP53|mutation|CISPLATIN"]["points"] == [
        {"sample_id": "S1", "group": "altered", "value": 4.5}
    ]


async def test_plots_route_streams_png_and_404s_unknown(tmp_cfg) -> None:
    sid = await _seed_report(tmp_cfg)
    async with _client(tmp_cfg) as c:
        r = await c.get(f"/api/sessions/{sid}/plots/kaplan_meier_1")
        assert r.status_code == 200
        assert r.headers["content-type"] == "image/png"
        assert r.content == PNG
        assert (await c.get(f"/api/sessions/{sid}/plots/nope")).status_code == 404
        other = await _make_session(tmp_cfg)
        assert (await c.get(f"/api/sessions/{other}/plots/kaplan_meier_1")).status_code == 404


async def test_plot_urls_are_encoded_and_still_route(tmp_cfg) -> None:
    sid = await _make_session(tmp_cfg)
    conn = await db_mod.connect(tmp_cfg)
    try:
        blob_id = await blobs.put(conn, "plots", "km.png", PNG, {"session_id": sid})
        await analysis_repo.save_report(conn, sid, {"plots": [
            {"name": "km #1?x", "blob_id": blob_id, "content_type": "image/png"}
        ]})
    finally:
        await conn.close()
    async with _client(tmp_cfg) as c:
        (plot,) = (await c.get(f"/api/sessions/{sid}/analysis")).json()["plots"]
        assert plot["url"] == f"/api/sessions/{sid}/plots/km%20%231%3Fx"
        r = await c.get(plot["url"])
    assert r.status_code == 200 and r.content == PNG


async def test_plots_route_404s_when_blob_is_gone(tmp_cfg) -> None:
    sid = await _make_session(tmp_cfg)
    conn = await db_mod.connect(tmp_cfg)
    try:
        await analysis_repo.save_report(conn, sid, {"plots": [
            {"name": "p", "blob_id": "not-an-object-id", "content_type": "image/png"}
        ]})
    finally:
        await conn.close()
    async with _client(tmp_cfg) as c:
        assert (await c.get(f"/api/sessions/{sid}/plots/p")).status_code == 404


async def test_attachments_route_404s_unknown_session(tmp_cfg) -> None:
    async with _client(tmp_cfg) as c:
        r = await c.get(f"/api/sessions/{ids.session_id()}/attachments")
    assert r.status_code == 404


# ------------------------------------------------------------- errors / ids --


def _session_routes(cfg: Config) -> list[tuple[str, str]]:
    """Every (method, path) under /api/sessions/{session_id}, as the app publishes them."""
    paths = create_app(cfg).openapi()["paths"]
    return [
        (method.upper(), path) for path, ops in paths.items()
        if path.startswith("/api/sessions/{session_id}") for method in ops
    ]


async def test_bad_session_id_is_400(tmp_cfg) -> None:
    routes = _session_routes(tmp_cfg)
    paths = {p for _, p in routes}
    for want in ("timeline", "research", "synthesis", "overview", "sources", "analysis",
                 "plots/{name}", "attachments", "events", "metrics", "feedback", "pause"):
        assert f"/api/sessions/{{session_id}}/{want}" in paths, want
    async with _client(tmp_cfg) as c:
        for bad in ("nope", "ses_short", "ses_" + "a" * 26, "ses_" + "A" * 27, "x%2E%2E"):
            for method, path in routes:
                url = path.replace("{session_id}", bad).replace("{name}", "n")
                r = await c.request(method, url)
                assert r.status_code == 400, (method, url, r.status_code)


async def test_real_session_ids_match_the_pattern() -> None:
    from co_scientist.web.app import SESSION_ID_RE

    assert all(SESSION_ID_RE.fullmatch(ids.session_id()) for _ in range(50))


async def test_pymongo_errors_become_503_without_detail(
    tmp_cfg, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from pymongo.errors import OperationFailure

    async def boom(conn, session_id):
        raise OperationFailure("secret host and query detail")

    import co_scientist.web.app as app_mod

    monkeypatch.setattr(app_mod.att_repo, "list_for_session", boom)
    sid = await _make_session(tmp_cfg)
    async with _client(tmp_cfg) as c:
        r = await c.get(f"/api/sessions/{sid}/attachments")
    assert r.status_code == 503
    assert "secret" not in r.text


# -------------------------------------------------------------- prune CLI --


async def test_prune_dry_run_then_yes(tmp_cfg, capsys) -> None:
    from types import SimpleNamespace

    from co_scientist.attachments.intake import IncomingFile
    from co_scientist.attachments.store import store_incoming
    from co_scientist.cli import attachments_prune

    sid = await _make_session(tmp_cfg)
    conn = await db_mod.connect(tmp_cfg)
    try:
        (meta,) = await store_incoming(conn, tmp_cfg, sid, [
            IncomingFile("old.csv", b"a,b\n", "text/csv", "tabular"),
        ])
    finally:
        await conn.close()
    ctx = SimpleNamespace(obj=(tmp_cfg, None))
    future = datetime.now(UTC) + timedelta(days=1)

    # Nothing is older than 30 days.
    await _run_sync(attachments_prune, ctx, 30, False)
    assert "No attachments" in capsys.readouterr().out

    # Make "older than 0 days" catch the file: it was uploaded before now.
    await _run_sync(attachments_prune, ctx, 0, False)
    out = capsys.readouterr().out
    assert "old.csv" in out and "dry run" in out.lower()
    conn = await db_mod.connect(tmp_cfg)
    try:
        assert await att_repo.fetch(conn, meta.id) is not None
    finally:
        await conn.close()

    await _run_sync(attachments_prune, ctx, 0, True)
    assert "Deleted 1" in capsys.readouterr().out
    conn = await db_mod.connect(tmp_cfg)
    try:
        assert await att_repo.fetch(conn, meta.id) is None
        with pytest.raises(KeyError):
            await blobs.get(conn, "attachments", meta.blob_id)
        assert await att_repo.older_than(conn, future) == []
    finally:
        await conn.close()


async def _run_sync(fn, *args):
    """Typer commands call asyncio.run; run them off the test's event loop."""
    import asyncio

    return await asyncio.to_thread(fn, *args)
